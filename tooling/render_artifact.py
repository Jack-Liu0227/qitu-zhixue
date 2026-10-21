"""Render an existing artifact's text using fixed code, never model-generated Python."""
from __future__ import annotations

import base64
import csv
import io
import json
import sys
import textwrap
import unicodedata

MIMES = {
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'pdf': 'application/pdf',
}


def wrap_slide_line(value: str) -> list[str]:
    result = []
    current = ''
    width = 0
    for character in value:
        step = 2 if unicodedata.east_asian_width(character) in ('W', 'F') else 1
        if width + step > 58:
            result.append(current)
            current = ''
            width = 0
        current += character
        width += step
    return result + [current]


def render(title: str, body: str, format_name: str) -> bytes:
    output = io.BytesIO()
    if format_name == 'docx':
        from docx import Document
        document = Document()
        document.add_heading(title, 0)
        for line in body.splitlines():
            document.add_paragraph(line)
        document.save(output)
        assert len(Document(io.BytesIO(output.getvalue())).paragraphs) > 0
    elif format_name == 'pptx':
        from pptx import Presentation
        from pptx.util import Pt
        document = Presentation()
        lines = [part for line in body.splitlines() for part in wrap_slide_line(line)]
        for start in range(0, max(1, len(lines)), 12):
            slide = document.slides.add_slide(document.slide_layouts[1])
            slide.shapes.title.text = title if start == 0 else f'{title} ({start // 12 + 1})'
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for index, line in enumerate(lines[start:start + 12]):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = line
                paragraph.font.size = Pt(18)
        document.save(output)
        assert len(Presentation(io.BytesIO(output.getvalue())).slides) > 0
    elif format_name == 'xlsx':
        from openpyxl import Workbook, load_workbook
        document = Workbook()
        sheet = document.active
        sheet.title = 'Project evidence'
        sheet.append([title])
        delimiter = '\t' if '\t' in body else ','
        for row in csv.reader(io.StringIO(body), delimiter=delimiter):
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        sheet.freeze_panes = 'A2'
        document.save(output)
        check = load_workbook(io.BytesIO(output.getvalue()))
        assert check.active['A1'].value == title
    elif format_name == 'pdf':
        import fitz
        document = fitz.open()
        lines = [part for line in (title + '\n\n' + body).splitlines() for part in (textwrap.wrap(line, 42) or [''])]
        for start in range(0, max(1, len(lines)), 42):
            page = document.new_page(width=595, height=842)
            for index, line in enumerate(lines[start:start + 42]):
                page.insert_text((42, 48 + index * 17), line, fontname='china-s', fontsize=11)
        document.save(output)
        with fitz.open(stream=output.getvalue(), filetype='pdf') as check:
            assert len(check) > 0
    else:
        raise ValueError('ARTIFACT_FORMAT_INVALID')
    return output.getvalue()


def main() -> None:
    try:
        value = json.load(sys.stdin)
        if set(value) != {'title', 'body', 'format'} or value['format'] not in MIMES:
            raise ValueError('ARTIFACT_FORMAT_INVALID')
        if not isinstance(value['title'], str) or not isinstance(value['body'], str) or not value['body'].strip():
            raise ValueError('ARTIFACT_TEXT_REQUIRED')
        result = render(value['title'], value['body'], value['format'])
        print(json.dumps({'ok': True, 'mimeType': MIMES[value['format']], 'base64': base64.b64encode(result).decode('ascii')}))
    except ImportError:
        print(json.dumps({'ok': False, 'code': 'ARTIFACT_RENDERER_DEPENDENCY_MISSING'}))
    except Exception:
        print(json.dumps({'ok': False, 'code': 'ARTIFACT_RENDER_FAILED'}))


if __name__ == '__main__':
    main()
