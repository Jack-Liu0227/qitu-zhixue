import io
import unittest

from render_artifact import render


class ArtifactRenderingTests(unittest.TestCase):
    def test_word_retains_learner_text(self):
        from docx import Document
        document = Document(io.BytesIO(render('探究记录', '模拟观察\n不能据此声称真实测量', 'docx')))
        self.assertIn('不能据此声称真实测量', [p.text for p in document.paragraphs])

    def test_slides_preserve_all_lines(self):
        from pptx import Presentation
        content = '\n'.join(f'Observation {i}' for i in range(50))
        document = Presentation(io.BytesIO(render('Evidence', content, 'pptx')))
        text = '\n'.join(shape.text for slide in document.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn('Observation 49', text)
        self.assertGreater(len(document.slides), 1)

    def test_workbook_never_executes_learner_formula_text(self):
        from openpyxl import load_workbook
        document = load_workbook(io.BytesIO(render('Data', 'value\n=1+1\n3', 'xlsx')))
        self.assertEqual(document.active['A3'].data_type, 's')
        self.assertEqual(document.active['A3'].value, '=1+1')

    def test_pdf_contains_chinese_and_late_content(self):
        import fitz
        content = '\n'.join(f'模型观察 {i}' for i in range(90))
        with fitz.open(stream=render('探究记录', content, 'pdf'), filetype='pdf') as document:
            text = ''.join(page.get_text() for page in document)
            self.assertIn('模型观察 89', text)
            self.assertGreater(len(document), 1)


if __name__ == '__main__':
    unittest.main()
